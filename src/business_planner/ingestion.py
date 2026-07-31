from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from pypdf import PdfReader

from .models import Chunk, Document

SUPPORTED = {".pdf", ".txt", ".md", ".csv", ".json", ".xlsx"}


def company_slug(name: str) -> str:
    slug = "".join(c.lower() if c.isalnum() else "_" for c in name.strip())
    return "_".join(part for part in slug.split("_") if part)


def _source_id(path: Path, page: int | None = None) -> str:
    raw = f"{path.as_posix()}#{page or 0}".encode()
    return f"src_{hashlib.sha256(raw).hexdigest()[:12]}"


def _category(relative_path: Path) -> str:
    return relative_path.parts[0] if len(relative_path.parts) > 1 else "uncategorized"


def _read_tabular(path: Path) -> str:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            lines.append(" | ".join("" if value is None else str(value) for value in row))
    return "\n".join(lines)


def load_documents(data_root: Path, company_name: str) -> list[Document]:
    company_dir = data_root / company_slug(company_name)
    if not company_dir.is_dir():
        raise FileNotFoundError(
            f"Company data directory not found: {company_dir}. "
            "Create it using the layout documented in README.md."
        )

    documents: list[Document] = []
    for path in sorted(company_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        relative = path.relative_to(company_dir)
        # Trusted agent configuration is not source evidence and must never be
        # retrieved by Reality or passed indirectly to Internal Audit.
        if relative.as_posix() == (
            "00_governance/executive_principles.json"
        ):
            continue
        category = _category(relative)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            for page_number, page in enumerate(PdfReader(path).pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    documents.append(
                        Document(_source_id(relative, page_number), company_name, category,
                                 relative, text, page_number)
                    )
        elif suffix in {".csv", ".xlsx"}:
            text = _read_tabular(path)
            documents.append(
                Document(_source_id(relative), company_name, category, relative, text)
            )
        else:
            text = path.read_text(encoding="utf-8-sig")
            if suffix == ".json":
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            documents.append(
                Document(_source_id(relative), company_name, category, relative, text)
            )
    if not documents:
        raise ValueError(f"No supported documents found under {company_dir}")
    return documents


def chunk_documents(
    documents: Iterable[Document], chunk_size: int = 1800, overlap: int = 200
) -> list[Chunk]:
    if chunk_size <= overlap or overlap < 0:
        raise ValueError("chunk_size must be greater than overlap >= 0")
    chunks: list[Chunk] = []
    for document in documents:
        normalized = " ".join(document.text.split())
        start = 0
        index = 0
        while start < len(normalized):
            end = min(start + chunk_size, len(normalized))
            if end < len(normalized):
                boundary = normalized.rfind(" ", start, end)
                if boundary > start + chunk_size // 2:
                    end = boundary
            text = normalized[start:end].strip()
            if text:
                chunks.append(
                    Chunk(document.source_id, document.company, document.category,
                          document.path, text, document.page, index)
                )
            if end >= len(normalized):
                break
            start = end - overlap
            index += 1
    return chunks
